"""The camera walk as a spreadsheet: <name>_cameras.xlsx (+ a CSV of the cameras).

    python tools/camera_table.py <scene name> [--out <folder>]

For the student package (pipeline\\requests\\004), item 3: one row per frame of
the video that the pipeline looked at, with where the phone was, where it
pointed, when, and where the sun stood -- in the package's one coordinate
system (site.json: Z up, Y north or the walk, origin on the ground below the
first camera, metres when scaled).

Sheets:
  Cameras          one row per extracted frame (placed, dropped as blurred, or
                   not placed by the solve)
  Scene            the capture: phone, date, place, solve, scale, north, up
  Sun on that day  every 30 minutes at the capture's place, sunrise and sunset
  Read me          what every column means

HONEST GAPS, written into the sheet rather than filled by a guess:
  - a phone video carries ONE GPS position for the whole clip, not one per
    frame. Per-camera latitude/longitude/altitude are worked out from it only
    when north AND scale are set (then: the first camera = that position);
    otherwise those cells say "needs north and scale".
  - heading is measured from the site's +Y axis; it is a compass bearing only
    once north is set.
  - clock times assume the phone's recorded creation time is the clip's start.

Place and time are personal data: they are written only with --with-place (for
a file that goes back to the person who filmed it -- Marc, 2026-09-27).
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

import colmap_cameras   # noqa: E402
import dense_mesh       # noqa: E402
import scene_paths      # noqa: E402
import sun_pos          # noqa: E402

M_PER_DEG_LAT = 111_320.0          # metres per degree of latitude (and of longitude at the equator)
DIAGONAL_35MM = 43.2666            # mm, the diagonal of a 36 x 24 mm frame


def frame_number(name: str) -> int:
    m = re.search(r"(\d+)", name)
    return int(m.group(1)) if m else -1


def parse_when(s: str | None) -> datetime | None:
    """"2026-09-11T11:12:21+0200" -> an aware datetime (None if unreadable)."""
    if not s:
        return None
    s = s.strip().replace("Z", "+00:00")
    s = re.sub(r"([+-]\d\d)(\d\d)$", r"\1:\2", s)
    try:
        d = datetime.fromisoformat(s)
    except ValueError:
        return None
    return d if d.tzinfo else None


def build(name: str, with_place: bool) -> dict:
    folder = scene_paths.scene_dir(name)
    rec = json.loads((folder / "capture.json").read_text(encoding="utf-8"))
    site = json.loads((folder / "site.json").read_text(encoding="utf-8"))
    sun = json.loads((folder / "sun.json").read_text(encoding="utf-8")) \
        if (folder / "sun.json").is_file() else {}
    work = scene_paths.ROOT / "work" / name
    model = dense_mesh.find_model(work)

    import source_meta
    meta = source_meta.read(Path(rec["source"])) if rec.get("source") else {}
    fps = float((meta.get("image") or {}).get("fps") or rec.get("fps") or 30.0)
    stride = int((rec.get("settings") or {}).get("stride") or 1)
    trim = float(((rec.get("settings") or {}).get("start")) or 0.0)
    start = parse_when(meta.get("when")) if with_place else None
    where = (meta.get("where") or {}) if with_place else {}

    M = np.asarray(site["matrix"], float)
    s = float(np.cbrt(np.linalg.det(M[:3, :3])))
    A = M[:3, :3] / s
    scaled = site["scale"]["m_per_unit"] is not None
    north = site["axes"]["y"] == "north"
    units = "m" if scaled else "scan units"

    cams = colmap_cameras.read_cameras_bin(model / "cameras.bin")
    placed = {}
    for im in colmap_cameras.read_images_bin(model / "images.bin"):
        R = np.asarray(colmap_cameras.quat_to_matrix(*im["q"]), float)
        c = -R.T @ np.asarray(im["t"], float)
        cam = cams[im["camera_id"]]
        placed[frame_number(im["name"])] = {
            "name": im["name"], "pos": (M @ np.append(c, 1))[:3],
            "right": A @ R[0], "down": A @ R[1], "fwd": A @ R[2], "cam": cam}
    on_disk = {frame_number(p.name) for p in (work / "images").glob("*.jpg")}
    last = max(set(placed) | on_disk)

    geo = None
    if with_place and scaled and north and where.get("lat") is not None and placed:
        first = placed[min(placed)]["pos"]
        geo = {"lat0": where["lat"], "lon0": where["lon"], "alt0": where.get("alt"), "p0": first}

    rows = []
    for k in range(1, last + 1):
        t_video = trim + (k - 1) * stride / fps
        row = {"frame": k, "video frame": (k - 1) * stride, "video time (s)": round(t_video, 3)}
        if start:
            row["date and time"] = (start + timedelta(seconds=t_video)).strftime("%Y-%m-%d %H:%M:%S")
        if k in placed:
            p = placed[k]
            f, r = p["fwd"], p["right"]
            fx, fy, cx, cy = p["cam"]["params"][:4]
            W, H = p["cam"]["width"], p["cam"]["height"]
            row.update({
                "status": "placed",
                "x": p["pos"][0], "y": p["pos"][1], "z": p["pos"][2],
                "height above origin": p["pos"][2],
                "heading (deg)": math.degrees(math.atan2(f[0], f[1])) % 360,
                "pitch (deg)": math.degrees(math.asin(max(-1, min(1, f[2])))),
                "roll (deg)": math.degrees(math.asin(max(-1, min(1, -r[2])))),
                "focal (px)": (fx + fy) / 2,
                "focal 35mm-equiv (mm)": (fx + fy) / 2 * DIAGONAL_35MM / math.hypot(W, H),
                "fov horizontal (deg)": math.degrees(2 * math.atan(W / (2 * fx))),
                "fov vertical (deg)": math.degrees(2 * math.atan(H / (2 * fy))),
                "image (px)": f"{W} x {H}",
            })
            if geo:
                d = p["pos"] - geo["p0"]
                lat = geo["lat0"] + d[1] / M_PER_DEG_LAT
                row["latitude"] = lat
                row["longitude"] = geo["lon0"] + d[0] / (M_PER_DEG_LAT * math.cos(math.radians(lat)))
                if geo["alt0"] is not None:
                    row["altitude (m)"] = geo["alt0"] + d[2]
            else:
                row["latitude"] = row["longitude"] = "needs north and scale" if with_place else "not included"
        else:
            row["status"] = "not placed by the solve" if k in on_disk else "dropped (blurred)"
        if start and where.get("lat") is not None:
            t = start + timedelta(seconds=t_video)
            sp = sun_pos.sun_position(sun_pos.utc_ms(t), where["lat"], where["lon"])
            row["sun elevation (deg)"] = sp["elevation"]
            row["sun azimuth (deg from true north)"] = sp["azimuth"]
            if north:
                az, el = math.radians(sp["azimuth"]), math.radians(sp["elevation"])
                row["sun direction x"] = math.sin(az) * math.cos(el)
                row["sun direction y"] = math.cos(az) * math.cos(el)
                row["sun direction z"] = math.sin(el)
        rows.append(row)

    day = []
    rise = set_ = None
    if start and where.get("lat") is not None:
        d0 = start.replace(hour=0, minute=0, second=0, microsecond=0)
        for i in range(48):
            t = d0 + timedelta(minutes=30 * i)
            sp = sun_pos.sun_position(sun_pos.utc_ms(t), where["lat"], where["lon"])
            day.append({"local time": t.strftime("%H:%M"),
                        "sun elevation (deg)": sp["elevation"],
                        "sun azimuth (deg from true north)": sp["azimuth"],
                        "above the horizon": "yes" if sp["elevation"] > 0 else "no"})
        rise, set_ = sun_pos.sunrise_sunset(start, where["lat"], where["lon"])

    n_placed = sum(1 for r in rows if r["status"] == "placed")
    scene = [
        ("scene", name),
        ("source video", Path(rec.get("source", "")).name),
        ("phone", " ".join(filter(None, [(meta.get("camera") or {}).get("make"),
                                         (meta.get("camera") or {}).get("model")]))),
        ("frames per second", round(fps, 3)),
        ("video length (s)", meta.get("duration_s")),
        ("video frame size (px)", f"{(meta.get('image') or {}).get('width')} x {(meta.get('image') or {}).get('height')}"),
        ("every Nth video frame used", stride),
        ("frames looked at", len(rows)),
        ("frames placed by the solve", n_placed),
        ("frames dropped as blurred", sum(1 for r in rows if r["status"] == "dropped (blurred)")),
        ("frames not placed", sum(1 for r in rows if r["status"] == "not placed by the solve")),
        ("coordinate system", "site.json: Z up, Y " + site["axes"]["y"] + ", X " + site["axes"]["x"]
         + "; origin on the ground below the first camera"),
        ("units", site["units"] + ("" if scaled else " -- NOT TO SCALE until a known distance is measured")),
        ("scale", site["scale"]["status"] + (f", {s:.6g} m per scan unit ({site['scale']['note']})" if scaled else "")),
        ("north", ("set: " + sun.get("north_source", "")) if north else "not set -- heading is measured from the walk direction, not from north"),
        ("up", site["how"]["z"]),
        ("origin", site["how"]["origin"]),
    ]
    if with_place:
        scene += [
            ("recorded (phone clock)", meta.get("when") or "not recorded"),
            ("GPS position (one for the whole clip)",
             f"{where.get('lat')}, {where.get('lon')}" if where.get("lat") is not None else "not recorded"),
            ("GPS altitude (m)", where.get("alt")),
            ("GPS accuracy (m)", where.get("accuracy_m")),
            ("sunrise (local)", rise.strftime("%H:%M") if rise else "none that day"),
            ("sunset (local)", set_.strftime("%H:%M") if set_ else "none that day"),
        ]
    scene += [("made by", "DL-SplatGenerator tools/camera_table.py (Digital Landscapes)"),
              ("made on", datetime.now().strftime("%Y-%m-%d %H:%M"))]
    return {"rows": rows, "day": day, "scene": scene, "units": units, "north": north,
            "scaled": scaled, "with_place": with_place, "name": name}


README_ROWS = [
    ("Cameras", ""),
    ("frame", "the pipeline's frame number (frame_00001 ...)"),
    ("video frame / video time (s)", "which frame of the video, and when in it (from the start of the clip)"),
    ("date and time", "clock time of that frame: the phone's recorded start time + the video time"),
    ("status", "placed = the solve found where the phone was; dropped (blurred) = too blurred to use; not placed = used but not solved"),
    ("x, y, z", "the camera's position in the package's coordinate system (see Scene: coordinate system, units)"),
    ("height above origin", "= z: the phone's height above the ground at the start of the walk"),
    ("heading (deg)", "direction the phone looked, clockwise from +Y (from north only once north is set)"),
    ("pitch (deg)", "up/down: negative = looking down"),
    ("roll (deg)", "sideways tilt: positive = the right edge of the picture lower"),
    ("focal, fov", "the phone's lens as the solve measured it, for the video's own frame (with its lens distortion); 35mm-equivalent by the diagonal. The cameras in the .3dm and .glb are the lens-corrected frame the splat and the mesh texture were made from, slightly narrower (a fraction of a degree to about one degree)"),
    ("latitude, longitude, altitude", "worked out from the clip's ONE GPS position (given to the first camera), only when north and scale are set"),
    ("sun elevation / azimuth", "where the sun stood at that moment (NOAA algorithm, the same as the viewer); azimuth clockwise from true north"),
    ("sun direction x, y, z", "a unit vector towards the sun in the package's coordinates -- only when north is set"),
    ("", ""),
    ("Sun on that day", "the sun every 30 minutes at the capture's place; sunrise/sunset in Scene"),
    ("", ""),
    ("Limits", "A reconstruction, not a survey. Until a known distance is measured the units are the scan's own, not metres. Phone GPS is typically good to a few metres; the recorded accuracy is in Scene."),
]


def write_xlsx(t: dict, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="30343A")

    def table(ws, rows: list[dict], formats: dict):
        cols = []
        for r in rows:
            for k in r:
                if k not in cols:
                    cols.append(k)
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in rows:
            ws.append([r.get(k) for k in cols])
        for i, k in enumerate(cols, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = max(10, min(28, len(k) + 2))
            fmt = next((f for key, f in formats.items() if key in k), None)
            if fmt:
                for cell in ws[letter][1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = fmt
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 45

    wb = Workbook()
    ws = wb.active
    ws.title = "Cameras"
    table(ws, t["rows"], {"(deg)": "0.00", "direction": "0.0000", "latitude": "0.000000",
                          "longitude": "0.000000", "altitude": "0.00", "focal": "0.0",
                          "fov": "0.00", "video time": "0.000", "x": "0.0000", "y": "0.0000",
                          "z": "0.0000", "height": "0.0000"})
    ws2 = wb.create_sheet("Scene")
    ws2.append(["", ""])
    for k, v in t["scene"]:
        ws2.append([k, v])
    ws2.delete_rows(1)
    ws2.column_dimensions["A"].width = 38
    ws2.column_dimensions["B"].width = 100
    for c in ws2["A"]:
        c.font = Font(bold=True)
    if t["day"]:
        ws3 = wb.create_sheet("Sun on that day")
        table(ws3, t["day"], {"(deg": "0.00"})
    ws4 = wb.create_sheet("Read me")
    ws4.append(["", ""])
    for k, v in README_ROWS:
        ws4.append([k, v])
    ws4.delete_rows(1)
    ws4.column_dimensions["A"].width = 32
    ws4.column_dimensions["B"].width = 110
    for c in ws4["A"]:
        c.font = Font(bold=True)
    wb.save(path)


def write_csv(t: dict, path: Path) -> None:
    cols = []
    for r in t["rows"]:
        for k in r:
            if k not in cols:
                cols.append(k)
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for r in t["rows"]:
            w.writerow([f"{v:.6f}" if isinstance(v, float) else ("" if v is None else v)
                        for v in (r.get(k) for k in cols)])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("name")
    ap.add_argument("--out", default=None,
                    help="folder (default: <scene>\\<name>\\<name> - site data\\cameras and sun\\)")
    ap.add_argument("--with-place", action="store_true",
                    help="include date, time, GPS and the sun (personal data)")
    args = ap.parse_args()
    folder = scene_paths.scene_dir(args.name)
    out = Path(args.out) if args.out else scene_paths.site_data_dir(args.name) / "cameras and sun"
    out.mkdir(parents=True, exist_ok=True)
    t = build(args.name, args.with_place)
    xlsx, csv_path = out / f"{args.name}_cameras.xlsx", out / f"{args.name}_cameras.csv"
    for p in (xlsx, csv_path):
        if p.exists():
            print(f"FAILED: {p} exists -- never overwritten; remove it or use --out", file=sys.stderr)
            return 2
    write_xlsx(t, xlsx)
    write_csv(t, csv_path)
    placed = sum(1 for r in t["rows"] if r["status"] == "placed")
    print(f"    {len(t['rows'])} frames ({placed} placed), units {t['units']}, "
          f"north {'set' if t['north'] else 'not set'}, place/time {'in' if t['with_place'] else 'left out'}")
    print(f"    written {xlsx}")
    print(f"    written {csv_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
